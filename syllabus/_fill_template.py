#!/usr/bin/env python3
# Fill the SJTU graduate-course syllabus template from the application-form content.
# Writes a NEW filled .docx (keeps the blank template intact).
import docx, copy
from docx.oxml.ns import qn

SRC = "syllabus/上海交通大学研究生课程教学大纲模版.docx"
OUT = "syllabus/量子科学计算-研究生课程教学大纲.docx"

doc = docx.Document(SRC)
t = doc.tables[0]

def _font(run):
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rf)
    rf.set(qn("w:eastAsia"), "宋体")

def setcell(cell, text):
    """Clear a (possibly merged) cell and write text; '\n' -> new paragraphs."""
    # remove all but first paragraph
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p0 = cell.paragraphs[0]
    for r in list(p0.runs):
        r._element.getparent().remove(r._element)
    lines = text.split("\n")
    r = p0.add_run(lines[0]); _font(r)
    for extra in lines[1:]:
        p = cell.add_paragraph()
        rr = p.add_run(extra); _font(rr)

def C(r, c):
    return t.rows[r].cells[c]

# --- r0: code / hours / credits ---
setcell(C(0,1), "")                     # 课程代码 Course Code (not specified in the form)
setcell(C(0,4), "48")                   # 学时 Teaching Hours
setcell(C(0,10), "3")                   # 学分 Credits
# --- course name (r1 Chinese, r2 English) ---
setcell(C(1,1), "量子科学计算")
setcell(C(2,1), "Quantum Scientific Computing")
# --- language / school / prerequisite ---
setcell(C(3,1), "中英文并行开班  Chinese & English")
setcell(C(4,1), "数学科学学院  School of Mathematical Sciences")
setcell(C(5,1), "线性代数、微积分、微分方程、数值方法与基础编程；量子力学非必须。\n"
                "Linear algebra, calculus, differential equations, numerical methods, and basic "
                "programming; quantum mechanics is not required.")
# --- instructor row r7: name / title / dept / email ---
setcell(C(7,1), "张镭  Lei Zhang")
setcell(C(7,3), "教授  Professor")
setcell(C(7,5), "数学科学学院与自然科学研究院\nSchool of Mathematical Sciences & Institute of Natural Sciences")
setcell(C(7,9), "lzhang2012@sjtu.edu.cn")

# --- course description (Chinese, ~350 chars) ---
desc_cn = ("本课程是一门面向科学计算应用的量子计算交叉课程，系统介绍量子计算的基本框架及其在科学与工程"
"计算中的应用。内容涵盖：量子态、量子门与量子线路模型；量子傅里叶变换、相位估计（QPE）、Hadamard "
"检验与振幅估计等核心算法原语；以 block encoding、线性组合幺正（LCU）、量子信号处理（QSP）、量子"
"奇异值变换（QSVT）与量子线性系统求解（HHL）为主线的量子数值线性代数；哈密顿量模拟、Trotter/LCU "
"方法与薛定谔化（Schrödingerization）等量子模拟技术；以及量子算法在偏微分方程、动力系统、金融、材料"
"和机器学习中的应用。课程目标是帮助学生掌握上述基本思想与方法，理解其在高维科学计算中的潜在优势"
"与局限，并借助 UnitaryLab 等平台进行上机实践。学生应具备线性代数、微积分、微分方程、数值方法和"
"基础编程知识；量子力学有助于理解部分内容，但并非必须先修课程。")
setcell(C(12,1), desc_cn)

desc_en = ("This course is an interdisciplinary quantum computing course oriented toward scientific "
"computing applications. It introduces quantum circuits, quantum algorithms, quantum linear algebra, "
"quantum simulation, and applications of quantum algorithms to partial differential equations, "
"dynamical systems, finance, materials science, and machine learning. The goal is to help students "
"understand the basic ideas and methods of quantum states, quantum gates, phase estimation, QSP, "
"QSVT, HHL, and Schrödingerization, and to recognize their potential advantages and limitations in "
"high-dimensional scientific computing. Students are expected to have background knowledge in linear "
"algebra, calculus, differential equations, numerical methods, and basic programming. Quantum "
"mechanics is helpful for understanding some parts of the course, but it is not a required prerequisite.")
setcell(C(13,1), desc_en)

# --- schedule rows r15..r21 (content c1, hours c6, format c8, instructor c11) ---
sched = [
 ("1. 量子计算基础与量子线路：量子态、量子门与测量；量子线路模型；普适性与 Clifford+T 综合；量子复杂度基本概念。\n"
  "Fundamentals of quantum computing & circuits: states, gates, measurement; circuit model; universality and Clifford+T synthesis; basics of quantum complexity.",
  "6", "讲授 + 推导 + 简单线路演示", "张镭"),
 ("2. 量子算法基础：QFT；相位估计（QPE）；Hadamard test；振幅估计；量子算法的频谱视角。\n"
  "Fundamentals of quantum algorithms: QFT; phase estimation; Hadamard test; amplitude estimation; spectral perspective.",
  "9", "讲授 + 算法推导 + 示例", "张镭"),
 ("3. 量子线性代数与后 QPE 方法：block encoding；量子信号处理（QSP）；量子奇异值变换（QSVT）；矩阵函数计算；量子线性系统（HHL）。\n"
  "Quantum linear algebra & post-QPE methods: block encoding; QSP; QSVT; matrix functions; quantum linear systems (HHL).",
  "9", "讲授 + 结构分析", "张镭"),
 ("4. 量子模拟与动力系统：Hamiltonian 模拟；Trotter/LCU 方法；Schrödingerization；特征值问题；谱方法与量子模拟的联系。\n"
  "Quantum simulation & dynamical systems: Hamiltonian simulation; Trotter/LCU; Schrödingerization; eigenvalue problems; spectral methods.",
  "9", "讲授 + 案例分析", "张镭"),
 ("5. 微分方程量子算法：线性与非线性 PDE（Poisson、波动方程、Hamilton–Jacobi 方程等）；薛定谔化方法；离散化与矩阵形式；误差来源（离散化/量子误差）；测量复杂性与可观测量计算。\n"
  "Quantum algorithms for differential equations: linear/nonlinear PDEs (Poisson, wave, Hamilton–Jacobi); Schrödingerization; discretization; error sources; measurement complexity.",
  "9", "讲授 + 文献分析", "张镭"),
 ("6. 量子科学计算应用专题（含实践）：金融（期权定价、风险分析）；材料（特征值问题、电子结构）；机器学习（核方法、回归）；高维问题与量子优势；UnitaryLab 与 Agent Coding 实践。\n"
  "Applications of quantum scientific computing (with practice): finance; materials; machine learning; high-dimensional problems and quantum advantage; UnitaryLab & agent-coding lab.",
  "6", "讲授 + 讨论 + 上机实验", "张镭"),
 ("合计  Total", "48", "", ""),
]
for k, (content, hours, fmt, inst) in enumerate(sched):
    r = 15 + k
    setcell(C(r,1), content)
    setcell(C(r,6), hours)
    setcell(C(r,8), fmt)
    setcell(C(r,11), inst)

# --- grading policy r22 ---
setcell(C(22,1),
 "课程考核采用过程性评价与综合评价相结合。平时作业 3–4 次（30%），考察基础知识与计算能力；"
 "小测验 1–2 次（20%），检验核心概念掌握情况；期末大作业（project，占 50%），完成一个综合性项目"
 "并提交报告。评分依据正确性、实现质量与表达情况；成绩以等级制记录。\n"
 "Assessment combines continuous and comprehensive evaluation: assignments (3–4, 30%); quizzes "
 "(1–2, 20%); and a final project with a written report (50%). Grading is based on correctness, "
 "implementation quality, and clarity of presentation; letter grading is used.")

# --- textbooks & references r23 ---
setcell(C(23,1),
 "自编讲义（Self-compiled lecture notes）。\n"
 "参考资料 / References:\n"
 "1. A. Childs, Lecture Notes on Quantum Algorithms, https://www.cs.umd.edu/~amchilds/qa/\n"
 "2. Lin Lin, Lecture Notes on Quantum Algorithms for Scientific Computation, https://math.berkeley.edu/~linlin/qasc/\n"
 "3. UnitaryLab, https://unitarylab.com/unitary-studio-test/")

# --- notes r24 ---
setcell(C(24,1),
 "本课程为硕博共用的专业选修课，秋季学期开设，中英文并行开班。\n"
 "A specialized elective for all graduate students; offered in Fall, taught bilingually (Chinese & English).")

doc.save(OUT)
print("saved:", OUT)
